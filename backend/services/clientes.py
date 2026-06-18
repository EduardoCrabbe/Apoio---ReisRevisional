"""Regras de negócio do CRM de clientes — Etapa 3.

Concentra a lógica (ownership, janela de 72h, transações de atendimento/quitação,
reset mensal, importação de planilha). Os routers ficam finos, só traduzindo
HTTP. Erros são levantados como HTTPException para propagar o status correto.
"""

import io
import json
import unicodedata
from datetime import datetime, timedelta, timezone

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

import models
from services.comissao import get_price

LIMITE_CONTATOS = 6
JANELA_HORAS = 72

# Cabeçalhos esperados na planilha de importação (campo -> nome exibido).
COLUNAS_IMPORT = {
    "id_datajuri": "Código DJ",
    "first_name": "Cliente",
    "uf": "UF",
    "contrato": "Tipo de contrato",
    "tem_processo": "Processo?",
}


# --------------------------------------------------------------------- helpers

def agora_utc() -> datetime:
    """Datetime UTC (naive), consistente com os defaults dos modelos."""
    return datetime.now(timezone.utc).replace(tzinfo=None)


def primeiro_nome(nome) -> str:
    """Apenas o primeiro nome (LGPD): 'João da Silva' -> 'João'."""
    texto = str(nome).strip() if nome is not None else ""
    return texto.split()[0] if texto else ""


def _normalizar(texto) -> str:
    s = str(texto if texto is not None else "").strip().lower()
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c))


def buscar_cliente(db: Session, customer_id: str) -> models.Customer:
    cliente = db.get(models.Customer, customer_id)
    if cliente is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Cliente não encontrado.")
    return cliente


def exigir_dono(cliente: models.Customer, user: models.User) -> None:
    """Ação de CS: o cliente PRECISA pertencer ao usuário logado (403 se alheio)."""
    if cliente.cs_id != user.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Este cliente não pertence a você.")


def exigir_pode_editar(cliente: models.Customer, user: models.User) -> None:
    """CS só edita os seus; Gerente/Supervisor editam qualquer um."""
    if user.role == "CS" and cliente.cs_id != user.id:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Este cliente não pertence a você.")


def resolver_cs_id(user: models.User, cs_id_informado):
    """cs_id = o próprio CS logado; para gestão é obrigatório informar."""
    if user.role == "CS":
        return user.id
    if cs_id_informado is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "cs_id é obrigatório para Gerente/Supervisor.")
    return cs_id_informado


def _exigir_nivel_cs(user: models.User) -> int:
    if user.level_cs is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Somente um CS com nível definido pode realizar esta ação.")
    return user.level_cs


# ----------------------------------------------------------------- consultas

def listar_clientes(db: Session, user: models.User, cs_id):
    q = db.query(models.Customer)
    if user.role == "CS":
        if cs_id is not None and cs_id != user.id:
            raise HTTPException(status.HTTP_403_FORBIDDEN, "Você só pode ver os seus clientes.")
        q = q.filter(models.Customer.cs_id == user.id)
    elif cs_id is not None:
        q = q.filter(models.Customer.cs_id == cs_id)
    return q.all()


def criar_cliente(db: Session, user: models.User, dados) -> models.Customer:
    cs_id = resolver_cs_id(user, dados.cs_id)
    if db.get(models.Customer, dados.id_datajuri) is not None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Já existe cliente com este Código DJ.")
    cliente = models.Customer(
        id_datajuri=str(dados.id_datajuri).strip(),
        first_name=primeiro_nome(dados.first_name),
        uf=dados.uf,
        contrato=dados.contrato,
        tem_processo=dados.tem_processo,
        criticidade=dados.criticidade,
        status="Ativo",
        cs_id=cs_id,
        contatos=0,
        tentativas=0,
    )
    db.add(cliente)
    db.commit()
    db.refresh(cliente)
    return cliente


# --------------------------------------------------------------- atendimento

def checar_atendimento_permitido(cliente: models.Customer) -> None:
    """Ordem das validações: limite de contatos ANTES da janela de 72h."""
    if cliente.contatos >= LIMITE_CONTATOS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Limite de 6 atendimentos atingido")
    if cliente.ultimo_contato is not None:
        restante = timedelta(hours=JANELA_HORAS) - (agora_utc() - cliente.ultimo_contato)
        segundos = int(restante.total_seconds())
        if segundos > 0:
            horas, minutos = segundos // 3600, (segundos % 3600) // 60
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                f"Próximo atendimento disponível em {horas}h {minutos}m",
            )


def registrar_atendimento(db: Session, cliente: models.Customer, user: models.User) -> models.Attendance:
    nivel = _exigir_nivel_cs(user)
    valor = get_price(nivel, "Atendimento", db)
    agora = agora_utc()
    att = models.Attendance(
        user_id=user.id,
        customer_id=cliente.id_datajuri,
        timestamp=agora,
        commission_value=valor,
    )
    db.add(att)
    cliente.contatos = (cliente.contatos or 0) + 1
    cliente.ultimo_contato = agora
    db.commit()  # transação única: attendance + contatos + ultimo_contato
    db.refresh(att)
    db.refresh(cliente)
    return att


def desfazer_atendimento(db: Session, cliente: models.Customer, user: models.User) -> None:
    if (cliente.contatos or 0) <= 0:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Não há atendimentos para desfazer.")

    def ultimo_do_par():
        return (
            db.query(models.Attendance)
            .filter(models.Attendance.user_id == user.id, models.Attendance.customer_id == cliente.id_datajuri)
            .order_by(models.Attendance.timestamp.desc(), models.Attendance.id.desc())
            .first()
        )

    ultimo = ultimo_do_par()
    if ultimo is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Não há atendimentos para desfazer.")
    db.delete(ultimo)
    db.flush()  # para a próxima consulta enxergar a remoção

    anterior = ultimo_do_par()
    cliente.contatos = max(0, (cliente.contatos or 0) - 1)
    # Restaura o timer para o atendimento anterior (ou null se era o único).
    cliente.ultimo_contato = anterior.timestamp if anterior else None
    db.commit()
    db.refresh(cliente)


# ------------------------------------------------------------------- quitação

def quitar(db: Session, cliente: models.Customer, user: models.User, dados):
    if dados.valor_pago > dados.valor_original:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "valor_pago não pode ser maior que valor_original.",
        )
    nivel = _exigir_nivel_cs(user)
    quitacao = models.Quitacao(
        customer_id=cliente.id_datajuri,
        cs_id=user.id,
        valor_original=dados.valor_original,
        valor_pago=dados.valor_pago,
        data_boleto=dados.data_boleto,
        data_pagamento=dados.data_pagamento,
        consulta_processo=dados.consulta_processo,
        protesto=dados.protesto,
        tarifas_restituiveis=dados.tarifas_restituiveis,
        mes_referencia=dados.mes_referencia or agora_utc().strftime("%Y-%m"),
    )
    db.add(quitacao)
    cliente.status = "Quitado"
    # Bônus de Quitação lançado AUTOMATICAMENTE (sem o CS lançar à parte).
    bonus = models.BonusEntry(
        user_id=user.id,
        customer_id=cliente.id_datajuri,
        tipo="Quitacao",
        valor=get_price(nivel, "Quitacao", db),
        timestamp=agora_utc(),
    )
    db.add(bonus)
    db.commit()
    db.refresh(quitacao)

    economia = dados.valor_original - dados.valor_pago
    percentual = round(economia / dados.valor_original, 4) if dados.valor_original else 0.0
    return quitacao, economia, percentual  # economia/percentual calculados, NUNCA salvos


# ---------------------------------------------------------------- reset mensal

def reset_mensal(db: Session, user: models.User) -> int:
    ativos = db.query(models.Customer).filter(models.Customer.status == "Ativo").all()
    for c in ativos:
        c.contatos = 0
        c.tentativas = 0
        c.ultimo_contato = None
    registro = json.dumps({
        "executado_por": user.id,
        "em": datetime.now(timezone.utc).isoformat(),
    })
    ss = db.query(models.SystemSettings).filter_by(key="ultimo_reset").first()
    if ss is None:
        db.add(models.SystemSettings(key="ultimo_reset", value=registro))
    else:
        ss.value = registro
    db.commit()
    return len(ativos)


# ----------------------------------------------------------------- importação

def _cell(row, idx):
    return row[idx] if idx is not None and idx < len(row) else None


def importar_clientes(db: Session, conteudo: bytes, cs_id) -> int:
    from openpyxl import load_workbook  # import lazy (dependência só da importação)

    try:
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Não foi possível ler o arquivo .xlsx.")

    ws = wb.active
    linhas = ws.iter_rows(values_only=True)
    try:
        cabecalho = next(linhas)
    except StopIteration:
        wb.close()
        raise HTTPException(422, "Planilha vazia.")

    norm_para_idx = {_normalizar(h): i for i, h in enumerate(cabecalho) if h is not None}

    idx, faltantes = {}, []
    for campo, nome in COLUNAS_IMPORT.items():
        pos = norm_para_idx.get(_normalizar(nome))
        if pos is None:
            faltantes.append(nome)
        else:
            idx[campo] = pos
    if faltantes:
        wb.close()
        raise HTTPException(
            422,
            f"Colunas faltantes: {', '.join(faltantes)}",
        )

    importados = 0
    for row in linhas:
        if row is None:
            continue
        codigo = _cell(row, idx["id_datajuri"])
        if codigo is None or str(codigo).strip() == "":
            continue
        codigo = str(codigo).strip()

        uf = _cell(row, idx["uf"])
        contrato = _cell(row, idx["contrato"])
        tem_processo = _cell(row, idx["tem_processo"])
        valores = dict(
            first_name=primeiro_nome(_cell(row, idx["first_name"])),
            uf=str(uf).strip() if uf is not None else None,
            contrato=str(contrato).strip() if contrato is not None else None,
            tem_processo=str(tem_processo).strip() if tem_processo is not None else "Não",
            cs_id=cs_id,
        )

        existente = db.get(models.Customer, codigo)
        if existente is not None:
            for k, v in valores.items():
                setattr(existente, k, v)
        else:
            db.add(models.Customer(
                id_datajuri=codigo,
                status="Ativo",
                criticidade="Regular",
                contatos=0,
                tentativas=0,
                **valores,
            ))
        importados += 1

    db.commit()
    wb.close()
    return importados
