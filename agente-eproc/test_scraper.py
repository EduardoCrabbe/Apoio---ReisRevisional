import os
import sys
from openpyxl import load_workbook
from scraper import EprocScraper

def run_tests(planilha_path=None):
    if planilha_path is None:
        if getattr(sys, 'frozen', False):
            base_dir = os.path.dirname(sys.executable)
            planilha_path = os.path.join(base_dir, "platform_manager", "base_clientes.xlsx")
        else:
            base_dir = os.path.dirname(__file__)
            planilha_path = os.path.join(base_dir, "..", "platform_manager", "base_clientes.xlsx")
    
    if not os.path.exists(planilha_path):
        print(f" Erro: Planilha no encontrada em {planilha_path}")
        return
        
    # Abriremos a planilha sem 'read_only' para podermos injetar os dados de volta
    wb = load_workbook(filename=planilha_path)
    sheet = wb.active
    
    # Preparando cabealhos para os resultados
    sheet.cell(row=1, column=3, value="Status Consulta")
    sheet.cell(row=1, column=4, value="Classe da Ao")
    sheet.cell(row=1, column=5, value="Data da Movimentao")
    sheet.cell(row=1, column=6, value="Descrio da Movimentao")
    sheet.cell(row=1, column=7, value="Triagem")
    
    
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) <= 1:
        print(" A planilha est vazia ou contm apenas o cabealho.")
        return
        
    dados = rows[1:]
    
    print(f" Inciando processamento em lote com {len(dados)} clientes da planilha.")
    
    resultados_consolidados = []
    
    # Rodando em modo normal (visvel) pois o Cloudflare exige janela no monitor
    with EprocScraper(headless=False) as scraper:
        for idx, linha in enumerate(dados):
            cpf = linha[0]
            processo_esperado = linha[1] if len(linha) > 1 else None
            
            if not cpf:
                continue
                
            print(f"\n==================================================")
            print(f"[{idx+1}/{len(dados)}]  Consultando CPF: {cpf}")
            if processo_esperado:
                print(f"Modo:  MONITORAMENTO (Processo Esperado: {processo_esperado})")
            else:
                print(f"Modo:  DESCOBERTA (Buscando processos de interesse)")
            print(f"==================================================")
            
            # Valor padro caso falhe
            resultado = {
                "Iterao": idx + 1,
                "CPF": cpf,
                "Processo": processo_esperado,
                "Status": "ERRO/NO CONSULTADO",
                "Triagem": "-"
            }
            
            try:
                # Agora passamos o processo esperado para que ele possa ser clicado na lista
                scraper.consultar_processo(cpf=cpf, processo_esperado=processo_esperado)
            except Exception as e:
                err_msg = str(e).encode('ascii', 'ignore').decode()
                print(f" Erro ao consultar {cpf}: {err_msg}")
                resultados_consolidados.append(resultado)
                sheet.cell(row=idx + 2, column=3, value="ERRO NA BUSCA")
                continue
            
            # O scraper salva o HTML atual no arquivo 'page_source.html'
            if os.path.exists("page_source.html"):
                with open("page_source.html", "r", encoding="utf-8") as f:
                    html = f.read()
                    
                    
                # BIFURCAO DA ARQUITETURA
                linha_planilha = idx + 2
                
                if processo_esperado:
                    # MODO MONITORAMENTO (Comportamento antigo)
                    if processo_esperado in html:
                        print(f" SUCESSO: Processo {processo_esperado} ENCONTRADO na pgina!")
                        
                        import re
                        
                        # Extrair Classe da Ao usando Regex direto no HTML
                        acao_str = ""
                        match_acao = re.search(r'Classe da a.o:.*?<[^>]+>([^<]+)', html, re.IGNORECASE | re.DOTALL)
                        if not match_acao:
                            # Fallback
                            match_acao = re.search(r'Classe da a.o:\s*([^<]+)', html, re.IGNORECASE)
                        
                        if match_acao:
                            acao_str = match_acao.group(1).strip()
                                
                        # Extrair a ltima movimentao (Data e Descrio) usando Regex
                        movimentacao_str = ""
                        data_str = ""
                        
                        # Procura a primeira TD que tenha uma data e captura ela e o contedo da TD logo em seguida
                        match_mov = re.search(r'<td[^>]*>\s*(\d{2}/\d{2}/\d{4} \d{2}:\d{2}:\d{2})\s*</td>\s*<td[^>]*>(.*?)</td>', html, re.IGNORECASE | re.DOTALL)
                        
                        if match_mov:
                            data_str = match_mov.group(1).strip()
                            # A descrio pode conter tags HTML dentro (como <a>), ento removemos as tags HTML internas para ficar limpo
                            desc_bruta = match_mov.group(2).strip()
                            movimentacao_str = re.sub(r'<[^>]+>', '', desc_bruta).strip()
                        
                        print(f" Classe da Ao Extrada: {acao_str if acao_str else 'NO ENCONTRADA'}")
                        print(f" ltima Movimentao (Data): {data_str if data_str else 'NO ENCONTRADA'}")
                        print(f" ltima Movimentao (Desc): {movimentacao_str if movimentacao_str else 'NO ENCONTRADA'}")
                        
                        alerta = "-"
                        if acao_str and movimentacao_str:
                            alerta = scraper.aplicar_triagem(acao_str, movimentacao_str)
                            print(f" Resultado da Triagem: {alerta}")
                        
                        resultado["Status"] = "ENCONTRADO"
                        resultado["Triagem"] = alerta
                        
                        # Salva os dados na planilha
                        sheet.cell(row=linha_planilha, column=3, value="ENCONTRADO")
                        sheet.cell(row=linha_planilha, column=4, value=acao_str if acao_str else "NO ENCONTRADA")
                        sheet.cell(row=linha_planilha, column=5, value=data_str if data_str else "NO ENCONTRADA")
                        sheet.cell(row=linha_planilha, column=6, value=movimentacao_str if movimentacao_str else "NO ENCONTRADA")
                        sheet.cell(row=linha_planilha, column=7, value=alerta)
                        
                    else:
                        print(f" FALHA: Processo {processo_esperado} NO encontrado na pgina.")
                        resultado["Status"] = "NO ENCONTRADO NA LISTA"
                        sheet.cell(row=linha_planilha, column=3, value="NO ENCONTRADO NA LISTA")
                        
                else:
                    # MODO DESCOBERTA (S tem o CPF)
                    html_upper = html.upper()
                    classes_perigosas = [
                        "BUSCA E APREENSO", "EXECUO DE TTULO EXTRAJUDICIAL",
                        "AO MONITRIA", "MONITRIA", "AO DE COBRANA", "COBRANA"
                    ]
                    
                    encontrou_perigo = False
                    for classe in classes_perigosas:
                        if classe in html_upper:
                            encontrou_perigo = True
                            print(f" DESCOBERTA: Possvel '{classe}' identificada na lista do CPF!")
                            break
                            
                    if encontrou_perigo:
                        resultado["Status"] = " DESCOBERTA: PROCESSO IDENTIFICADO"
                        resultado["Triagem"] = " REQUER VALIDAO DO CS"
                        
                        sheet.cell(row=linha_planilha, column=3, value="PROCESSO IDENTIFICADO")
                        sheet.cell(row=linha_planilha, column=7, value=" REQUER VALIDAO DO CS")
                    else:
                        print(f" DESCOBERTA: Nenhum processo perigoso encontrado na lista para este CPF.")
                        resultado["Status"] = "NENHUM PROCESSO ENCONTRADO"
                        resultado["Triagem"] = " LIMPO"
                        
                        sheet.cell(row=linha_planilha, column=3, value="NENHUM PROCESSO ENCONTRADO")
                        sheet.cell(row=linha_planilha, column=7, value=" LIMPO")
                        
            else:
                print(" FALHA: Arquivo page_source.html no foi gerado.")
                sheet.cell(row=idx + 2, column=3, value="ERRO HTML")
                
            resultados_consolidados.append(resultado)

    # Imprimir o Relatrio Consolidado no final
    print("\n\n" + "="*80)
    print(" RELATRIO CONSOLIDADO DAS CONSULTAS")
    print("="*80)
    print(f"{'#':<4} | {'CPF':<15} | {'PROCESSO':<26} | {'STATUS':<22} | {'TRIAGEM'}")
    print("-"*80)
    for res in resultados_consolidados:
        print(f"{res['Iterao']:<4} | {res['CPF']:<15} | {res['Processo']:<26} | {res['Status']:<22} | {res['Triagem']}")
    print("="*80)
    
    # Salvar a planilha com os dados atualizados
    try:
        wb.save(planilha_path)
        print(f"\n Planilha salva com sucesso com as informaes extradas!")
    except PermissionError:
        print(f"\n ERRO FATAL: O rob no conseguiu salvar a planilha porque ela est aberta no Excel. Por favor, feche a planilha e rode o rob novamente.")

if __name__ == "__main__":
    run_tests()
