"""Agente Eproc — modelo DESACOPLADO (local-primeiro, sincroniza depois).

Roda na máquina do CS. Importa a planilha do dia, consulta o Eproc (real ou
simulado), grava o resultado LOCALMENTE e tenta enviar ao servidor em
best-effort. O CPF e o número do processo NUNCA saem desta máquina: só os campos
seguros {id_datajuri, classe, data_movimentacao, descricao, triagem} são enviados.

O agente NÃO depende do servidor para funcionar — se o servidor cair, o trabalho
do CS continua e os envios ficam pendentes para o próximo ciclo.

Uso:
    python agente.py                 # roda o dia (importa planilha, consulta, grava, envia)
    python agente.py encerrar-dia    # apaga planilha importada e resultados locais

Config: config.json (ao lado do agente) ou variáveis de ambiente:
    SERVER_URL, ROBO_TOKEN, PLANILHA_DIA, MODO_SIMULADO (0/1)
"""

import argparse
import json
import logging
import os
import sqlite3
import sys
import unicodedata
import urllib.error
import urllib.request
from datetime import date, datetime, timedelta, timezone

# Base de arquivos: ao lado do .py (ou do .exe quando empacotado com PyInstaller).
if getattr(sys, "frozen", False):
    BASE_DIR = os.path.dirname(sys.executable)
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))

DB_LOCAL = os.path.join(BASE_DIR, "resultados_local.db")
HTML_LOCAL = os.path.join(BASE_DIR, "relatorio_dia.html")
PENDENTES = os.path.join(BASE_DIR, "pendentes_local.json")
LOG_FILE = os.path.join(BASE_DIR, "agente.log")
CONFIG_FILE = os.path.join(BASE_DIR, "config.json")

ALERTA_VERMELHO = "🚨 ALERTA VERMELHO"
CAMPOS_SEGUROS = ("id_datajuri", "classe", "data_movimentacao", "descricao", "triagem")

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.FileHandler(LOG_FILE, encoding="utf-8"), logging.StreamHandler()],
)
log = logging.getLogger("agente-eproc")


# --------------------------------------------------------------------- config

def carregar_config() -> dict:
    cfg = {}
    if os.path.exists(CONFIG_FILE):
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            cfg = json.load(f)
    return {
        "SERVER_URL": cfg.get("SERVER_URL") or os.getenv("SERVER_URL", "http://localhost:8000"),
        "ROBO_TOKEN": cfg.get("ROBO_TOKEN") or os.getenv("ROBO_TOKEN", ""),
        "PLANILHA_DIA": cfg.get("PLANILHA_DIA") or os.getenv("PLANILHA_DIA", os.path.join(BASE_DIR, "planilha_exemplo.xlsx")),
        "MODO_SIMULADO": str(cfg.get("MODO_SIMULADO", os.getenv("MODO_SIMULADO", "1"))) == "1",
    }


# --------------------------------------------------------------------- helpers

def _norm(texto) -> str:
    s = str(texto if texto is not None else "").upper()
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c))


def triagem(classe, movimentacao) -> str:
    """classe contém 'BUSCA E APREENSÃO' E movimentação contém 'Petição'/'Mandado'."""
    c, m = _norm(classe), _norm(movimentacao)
    if "BUSCA E APREENSAO" in c and ("PETICAO" in m or "MANDADO" in m):
        return ALERTA_VERMELHO
    return "NORMAL"


def ler_planilha(caminho: str):
    """Lê id_datajuri, cpf, processo (cabeçalhos tolerantes a acento/caixa)."""
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.active
    linhas = ws.iter_rows(values_only=True)
    cabecalho = next(linhas)
    idx = {_norm(h): i for i, h in enumerate(cabecalho) if h is not None}

    def pos(*nomes):
        for n in nomes:
            if _norm(n) in idx:
                return idx[_norm(n)]
        return None

    i_id, i_cpf, i_proc = pos("id_datajuri", "codigo dj"), pos("cpf"), pos("processo", "numero do processo")
    if i_id is None:
        wb.close()
        raise ValueError("Planilha sem a coluna 'id_datajuri'.")

    registros = []
    for row in linhas:
        if row is None or i_id >= len(row) or row[i_id] in (None, ""):
            continue
        registros.append({
            "id_datajuri": str(row[i_id]).strip(),
            "cpf": str(row[i_cpf]).strip() if i_cpf is not None and i_cpf < len(row) and row[i_cpf] else "",
            "processo": str(row[i_proc]).strip() if i_proc is not None and i_proc < len(row) and row[i_proc] else "",
        })
    wb.close()
    return registros


# ---------------------------------------------------------------- consulta

def consultar_simulado(i: int, registro: dict) -> dict:
    """Resultado fictício realista, SEM rede. Garante 1 ALERTA VERMELHO (i==0)."""
    hoje = date.today()
    if i == 0:
        return {
            "classe": "BUSCA E APREENSÃO",
            "data_movimentacao": hoje.isoformat(),
            "descricao": "Juntada de Petição requerendo o cumprimento do Mandado de busca e apreensão do veículo.",
        }
    opcoes = [
        ("Procedimento Comum Cível", "Juntada de petição de manifestação da parte autora."),
        ("Procedimento Comum Cível", "Despacho: aguarde-se o decurso de prazo."),
        ("Cumprimento de Sentença", "Conclusos para decisão."),
        ("Procedimento Comum Cível", "Audiência de conciliação designada."),
    ]
    classe, desc = opcoes[i % len(opcoes)]
    return {"classe": classe, "data_movimentacao": (hoje - timedelta(days=i)).isoformat(), "descricao": desc}


def consultar_real(registro: dict) -> dict:
    """Ponto de integração com o scraper EXISTENTE (não alterar a lógica dele).

    Requer undetected-chromedriver etc. Ajuste o parse conforme o retorno de
    consultar_processo() do seu scraper. O demo/testes usam MODO_SIMULADO=1.
    """
    try:
        from scraper import EprocScraper  # scraper existente em agente-eproc/
    except Exception as exc:  # pragma: no cover
        raise RuntimeError(f"Scraper real indisponível: {exc}")

    with EprocScraper(headless=True) as s:  # pragma: no cover
        bruto = s.consultar_processo(numero_processo=registro["processo"], cpf=registro["cpf"])
        # 'bruto' depende do scraper; tratamos como descrição da movimentação.
        return {"classe": "", "data_movimentacao": date.today().isoformat(), "descricao": str(bruto)}


# --------------------------------------------------------- armazenamento local

def _conn():
    con = sqlite3.connect(DB_LOCAL)
    con.execute("""
        CREATE TABLE IF NOT EXISTS resultados (
            id_datajuri TEXT, cpf TEXT, processo TEXT, classe TEXT,
            data_movimentacao TEXT, descricao TEXT, triagem TEXT, registrado_em TEXT
        )""")
    return con


def gravar_local(registro: dict, res: dict, triagem_str: str):
    """Grava TUDO localmente (inclui cpf/processo — nunca sai daqui)."""
    con = _conn()
    con.execute(
        "INSERT INTO resultados VALUES (?,?,?,?,?,?,?,?)",
        (registro["id_datajuri"], registro["cpf"], registro["processo"], res["classe"],
         res["data_movimentacao"], res["descricao"], triagem_str,
         datetime.now(timezone.utc).isoformat()),
    )
    con.commit()
    con.close()


def gerar_html():
    """Relatório local para o CS ver o andamento (SEM CPF, por segurança)."""
    con = _conn()
    linhas = con.execute(
        "SELECT id_datajuri, classe, data_movimentacao, descricao, triagem FROM resultados ORDER BY registrado_em"
    ).fetchall()
    con.close()
    cabecalho = "<tr><th>id_datajuri</th><th>Classe</th><th>Data</th><th>Movimentação</th><th>Triagem</th></tr>"
    corpo = ""
    for r in linhas:
        cor = "#fde2e1" if r[4] == ALERTA_VERMELHO else "#ffffff"
        corpo += f"<tr style='background:{cor}'>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>"
    html = (
        "<html><head><meta charset='utf-8'><title>Relatório do dia — Eproc</title>"
        "<style>body{font-family:sans-serif}table{border-collapse:collapse;width:100%}"
        "td,th{border:1px solid #ccc;padding:6px;text-align:left}</style></head>"
        f"<body><h2>Relatório do dia — Eproc ({date.today().isoformat()})</h2>"
        f"<p>Sem CPF/processo — esses dados ficam só no banco local.</p>"
        f"<table>{cabecalho}{corpo}</table></body></html>"
    )
    with open(HTML_LOCAL, "w", encoding="utf-8") as f:
        f.write(html)


# ------------------------------------------------------------ envio best-effort

def _payload_seguro(registro: dict, res: dict, triagem_str: str) -> dict:
    p = {"id_datajuri": registro["id_datajuri"], "classe": res["classe"],
         "data_movimentacao": res["data_movimentacao"], "descricao": res["descricao"],
         "triagem": triagem_str}
    # Cinto de segurança: jamais incluir cpf/processo/nome.
    return {k: p[k] for k in CAMPOS_SEGUROS}


def enviar_servidor(server_url: str, token: str, payload: dict, timeout: int = 10) -> bool:
    url = server_url.rstrip("/") + "/api/robo/resultado"
    dados = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        url, data=dados, method="POST",
        headers={"Content-Type": "application/json", "X-Robo-Token": token},
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return 200 <= resp.status < 300
    except Exception as exc:
        log.warning("Servidor indisponível para %s: %s", payload["id_datajuri"], exc)
        return False


def _carregar_pendentes():
    if os.path.exists(PENDENTES):
        with open(PENDENTES, "r", encoding="utf-8") as f:
            return json.load(f)
    return []


def _salvar_pendentes(lista):
    with open(PENDENTES, "w", encoding="utf-8") as f:
        json.dump(lista, f, ensure_ascii=False, indent=2)


def reenviar_pendentes(cfg):
    pendentes = _carregar_pendentes()
    if not pendentes:
        return
    log.info("Reenviando %d pendente(s)...", len(pendentes))
    restantes = [p for p in pendentes if not enviar_servidor(cfg["SERVER_URL"], cfg["ROBO_TOKEN"], p)]
    _salvar_pendentes(restantes)
    log.info("Pendentes restantes: %d", len(restantes))


# ----------------------------------------------------------------- comandos

def rodar_dia():
    cfg = carregar_config()
    log.info("Iniciando ciclo. MODO_SIMULADO=%s | servidor=%s", cfg["MODO_SIMULADO"], cfg["SERVER_URL"])

    reenviar_pendentes(cfg)  # tenta limpar a fila local antes de começar

    registros = ler_planilha(cfg["PLANILHA_DIA"])
    log.info("Planilha importada: %d cliente(s).", len(registros))

    novos_pendentes = _carregar_pendentes()
    for i, registro in enumerate(registros):
        try:
            res = consultar_simulado(i, registro) if cfg["MODO_SIMULADO"] else consultar_real(registro)
        except Exception as exc:
            log.error("Falha ao consultar %s: %s", registro["id_datajuri"], exc)
            continue

        tri = triagem(res["classe"], res["descricao"])
        gravar_local(registro, res, tri)  # local-primeiro: sempre grava
        log.info("%s -> %s", registro["id_datajuri"], tri)

        payload = _payload_seguro(registro, res, tri)
        if not enviar_servidor(cfg["SERVER_URL"], cfg["ROBO_TOKEN"], payload):
            novos_pendentes.append(payload)  # best-effort: enfileira p/ depois

    _salvar_pendentes(novos_pendentes)
    gerar_html()
    log.info("Ciclo concluído. Relatório local: %s", HTML_LOCAL)


def encerrar_dia():
    for caminho in (DB_LOCAL, HTML_LOCAL, PENDENTES):
        if os.path.exists(caminho):
            os.remove(caminho)
            log.info("Removido: %s", caminho)
    cfg = carregar_config()
    planilha = cfg["PLANILHA_DIA"]
    if os.path.exists(planilha):
        os.remove(planilha)
        log.info("Planilha do dia removida: %s", planilha)
    log.info("Dia encerrado. Dados locais apagados.")


def main():
    parser = argparse.ArgumentParser(description="Agente Eproc (local-primeiro)")
    parser.add_argument("comando", nargs="?", default="rodar", choices=["rodar", "encerrar-dia"])
    args = parser.parse_args()
    if args.comando == "encerrar-dia":
        encerrar_dia()
    else:
        rodar_dia()


if __name__ == "__main__":
    main()
